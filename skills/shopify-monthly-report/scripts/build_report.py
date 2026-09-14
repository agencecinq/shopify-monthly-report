#!/usr/bin/env python3
"""
Builds the monthly deck from a report.json and the store's Shopify theme.

    python3 build_report.py --data report.json \
        --theme settings_data.json \
        --logo logo.png \
        --config client.json \
        --state plan.json \
        --lang fr \
        -o report.pptx

The deck adapts: any section missing from report.json is simply skipped, and
no empty slide is ever produced. Every user-visible string comes from
i18n/<lang>.json, never from this file.
"""

import argparse
import json
import os
import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from brand import build_palette, contrast, mix  # noqa: E402
from i18n import Lang, available  # noqa: E402

W, H = 13.333, 7.5
M = 0.62               # marge laterale
TOP = 0.55             # haut du titre
BODY_TOP = 1.72        # haut de la zone de contenu
FOOT = 6.95

NBSP = " "


# ---------------------------------------------------------------------------
# formatage
# ---------------------------------------------------------------------------

# One language object per run. A CLI process builds exactly one deck, so a
# module-level handle keeps every call site free of plumbing.
_L = Lang("en")


def set_lang(lang):
    global _L
    _L = lang


def T(key, **kwargs):
    return _L.t(key, **kwargs)


def money(v, cur="EUR", short=False):
    return _L.money(v, cur, short)


def integer(v):
    return _L.num(v)


def percent(v, dec=1):
    return _L.pct(v, dec)


def fmt(value, kind, cur="EUR"):
    if kind == "money":
        return money(value, cur, short=True)
    if kind == "percent":
        return percent(value)
    if kind == "decimal":
        return _L.num(value, 2) if value is not None else T("common.na")
    return integer(value)


def delta(cur_v, ref_v, magnitude=False):
    """
    Relative change. Returns (text, direction) or (None, 0) when it cannot be
    computed.

    `magnitude=True` compares absolute values. This is required for the lines
    Shopify stores as negatives (discounts, refunds): without it, 6,517 of
    refunds against 7,417 the month before reads as +12.1%, when there were
    fewer refunds.
    """
    if cur_v is None or ref_v in (None, 0):
        return None, 0
    a, b = (abs(cur_v), abs(ref_v)) if magnitude else (cur_v, ref_v)
    if b == 0:
        return None, 0
    d = (a - b) / abs(b)
    sign = "+" if d >= 0 else ""
    return f"{sign}{_L.num(d * 100, 1)}{_L.percent_space}%", (
        1 if d > 0 else (-1 if d < 0 else 0))


# ---------------------------------------------------------------------------
# primitives de mise en page
# ---------------------------------------------------------------------------

class Deck:
    def __init__(self, palette, meta):
        self.p = palette
        self.meta = meta
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.page = 0
        self.chapter = None
        self.toc = []
        self.money_fmt = _L.chart_number_format(meta.get("currency", "EUR"))

    # -- helpers bas niveau ------------------------------------------------
    def rgb(self, key):
        return RGBColor.from_string(self.p[key] if key in self.p else key)

    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = self.rgb("invert_bg" if dark else "bg")
        return s

    def text(self, s, x, y, w, h, runs, size=14, bold=False, color=None,
             font=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, wrap=True, shrink=False):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        if shrink:
            tf.auto_size = None
        items = runs if isinstance(runs, list) else [{"t": str(runs)}]
        first = True
        for item in items:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = item.get("align", align)
            if spacing:
                para.line_spacing = spacing
            if item.get("space_before"):
                para.space_before = Pt(item["space_before"])
            r = para.add_run()
            r.text = item["t"]
            f = r.font
            f.size = Pt(item.get("size", size))
            f.bold = item.get("bold", bold)
            f.italic = item.get("italic", False)
            f.name = item.get("font", font or self.p["font_body"])
            f.color.rgb = RGBColor.from_string(item.get("color", color or self.p["ink"]))
        return box

    def card(self, s, x, y, w, h, fill=None, radius=0.035):
        shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
        shp.adjustments[0] = radius
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill or self.p["surface"])
        shp.line.fill.background()
        _flatten(shp)
        if shp.text_frame:
            shp.text_frame.text = ""
        return shp

    def bar(self, s, x, y, w, h, fill, radius=None):
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            shp.adjustments[0] = radius
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(fill)
        shp.line.fill.background()
        _flatten(shp)
        return shp

    # -- chrome de slide ---------------------------------------------------
    def frame(self, s, title, kicker=None, dark=False, hint=None):
        """
        Pose le titre, la base de comparaison et le pied de slide.
        Retourne l'ordonnée à laquelle le contenu peut commencer : la ligne de
        lecture pédagogique, quand elle existe, décale tout le corps.
        """
        ink = self.p["invert_ink"] if dark else self.p["ink"]
        muted = self.p["invert_muted"] if dark else self.p["muted"]
        self.text(s, M, TOP, W - 2 * M - 3.6, 0.62, title, size=27, bold=True,
                  color=ink, font=self.p["font_heading"], anchor=MSO_ANCHOR.MIDDLE)
        if kicker:
            self.text(s, W - M - 3.6, TOP, 3.6, 0.62, kicker, size=10.5,
                      color=muted, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        y = BODY_TOP
        if hint:
            self.text(s, M, TOP + 0.60, W - 2 * M, 0.34, hint, size=10.5,
                      color=muted, anchor=MSO_ANCHOR.MIDDLE)
            y = BODY_TOP + 0.30
        self.footer(s, dark=dark)
        return y

    def footer(self, s, dark=False):
        self.page += 1
        muted = self.p["invert_muted"] if dark else self.p["muted"]
        left = f"{self.meta['shop_name']} · {self.meta['period_label']}"
        if self.chapter:
            left += f" · {self.chapter}"
        left += " · " + T("footer.source")
        agency = self.meta.get("agency")
        if agency:
            left += " · " + T("footer.agency", agency=agency)
        self.text(s, M, FOOT, W - 2 * M - 0.8, 0.3, left, size=8, color=muted,
                  anchor=MSO_ANCHOR.MIDDLE)
        self.text(s, W - M - 0.7, FOOT, 0.7, 0.3, str(self.page), size=8,
                  color=muted, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # -- graphiques --------------------------------------------------------
    def _chart_chrome(self, chart, legend=False, value_fmt='#,##0',
                      cat_rotation=None):
        p = self.p
        chart.font.size = Pt(9.5)
        chart.font.name = p["font_body"]
        chart.font.color.rgb = self.rgb("muted")
        chart.has_title = False
        chart.has_legend = legend
        if legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(9.5)
            chart.legend.font.color.rgb = self.rgb("muted")
        # fond transparent, pas de cadre
        el = chart._chartSpace
        sp = el.find(qn("c:spPr"))
        if sp is None:
            sp = el.makeelement(qn("c:spPr"), {})
            el.insert(0, sp)
        sp.clear()
        sp.append(sp.makeelement(qn("a:noFill"), {}))
        ln = sp.makeelement(qn("a:ln"), {})
        ln.append(ln.makeelement(qn("a:noFill"), {}))
        sp.append(ln)
        try:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = self.rgb("hairline")
            va.major_gridlines.format.line.width = Pt(0.75)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(9)
            va.tick_labels.font.color.rgb = self.rgb("muted")
            va.tick_labels.number_format = value_fmt
            va.tick_labels.number_format_is_linked = False
            va.has_title = False
        except (ValueError, NotImplementedError):
            pass
        try:
            ca = chart.category_axis
            ca.has_major_gridlines = False
            ca.format.line.color.rgb = self.rgb("hairline")
            ca.tick_labels.font.size = Pt(9)
            ca.tick_labels.font.color.rgb = self.rgb("muted")
            if cat_rotation is not None:
                ca.tick_labels.rotation = cat_rotation
            ca.has_title = False
        except (ValueError, NotImplementedError):
            pass

    def line_chart(self, s, x, y, w, h, categories, series, value_fmt='#,##0'):
        cd = CategoryChartData()
        cd.categories = categories
        for name, values in series:
            cd.add_series(name, values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
        chart = gf.chart
        self._chart_chrome(chart, legend=len(series) > 1, value_fmt=value_fmt)
        for i, plot_series in enumerate(chart.series):
            col = self.p["series"][i % len(self.p["series"])]
            line = plot_series.format.line
            line.color.rgb = RGBColor.from_string(col)
            line.width = Pt(2.5 if i == 0 else 1.5)
            plot_series.smooth = False
            if i > 0:
                _dash(plot_series, "dash")
        return chart

    def bar_chart(self, s, x, y, w, h, categories, values, horizontal=False,
                  value_fmt='#,##0', labels=True, colors=None, gap=60):
        cd = CategoryChartData()
        cd.categories = categories
        cd.add_series("s", values)
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = s.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
        chart = gf.chart
        self._chart_chrome(chart, legend=False, value_fmt=value_fmt)
        plot = chart.plots[0]
        plot.gap_width = gap
        plot.vary_by_categories = bool(colors)
        ser = chart.series[0]
        if colors:
            for idx, pt in enumerate(ser.points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = RGBColor.from_string(
                    colors[idx % len(colors)])
        else:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGBColor.from_string(self.p["series"][0])
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.number_format = value_fmt
            dl.number_format_is_linked = False
            dl.font.size = Pt(9)
            dl.font.bold = True
            dl.font.name = self.p["font_body"]
            dl.font.color.rgb = self.rgb("ink")
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        return chart

    def doughnut(self, s, x, y, w, h, categories, values):
        cd = CategoryChartData()
        cd.categories = categories
        cd.add_series("s", values)
        gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
        chart = gf.chart
        self._chart_chrome(chart, legend=True)
        plot = chart.plots[0]
        for idx, pt in enumerate(chart.series[0].points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(
                self.p["series"][idx % len(self.p["series"])])
            pt.format.line.color.rgb = self.rgb("bg")
            pt.format.line.width = Pt(2)
        plot.has_data_labels = False
        return chart


def _flatten(shape):
    """
    Supprime toute ombre portee. `shadow.inherit = False` ne suffit pas: les
    formes creees par python-pptx portent un bloc <p:style> qui pointe vers
    l'effet du theme, et certains lecteurs le rendent quand meme.
    """
    sp = shape._element
    for style in sp.findall(qn("p:style")):
        sp.remove(style)
    spPr = sp.find(qn("p:spPr"))
    if spPr is not None:
        for old in spPr.findall(qn("a:effectLst")):
            spPr.remove(old)
        spPr.append(spPr.makeelement(qn("a:effectLst"), {}))


def _dash(series, style="dash"):
    ln = series.format.line._get_or_add_ln()
    for tag in ("a:prstDash",):
        for old in ln.findall(qn(tag)):
            ln.remove(old)
    el = ln.makeelement(qn("a:prstDash"), {"val": style})
    ln.append(el)



# ---------------------------------------------------------------------------
# reading lines
# ---------------------------------------------------------------------------
# One sentence per chapter, shown under the title, saying what the slide
# measures and how to read it. Lives in the language catalogue under `hints.`,
# and report.json may override any of them through `lectures`.

def lecture(data, key):
    override = (data.get("lectures") or {}).get(key)
    if override:
        return override
    return T(f"hints.{key}") if _L.has(f"hints.{key}") else None


# ---------------------------------------------------------------------------
# slides
# ---------------------------------------------------------------------------

def _place_logo(d, s, logo, x, y, bg_hex, height=0.72):
    """
    Pose le logo, et glisse une pastille claire dessous quand le logo est trop
    sombre pour le fond. Sans cette précaution un logo noir disparaît sur une
    couverture noire, ce qui arrive dès que la marque est monochrome.
    """
    try:
        from PIL import Image
        img = Image.open(logo).convert("RGBA")
        iw, ih = img.size
        w_in = max(0.4, min(3.4, height * iw / max(1, ih)))
        small = img.resize((32, 32))
        px = [q for q in small.getdata() if q[3] > 40]
        if px:
            avg = sum(0.2126 * r + 0.7152 * g + 0.0722 * b for r, g, b, _ in px) / len(px)
            lum_logo = avg / 255.0
            lum_bg = (0.2126 * int(bg_hex[0:2], 16) + 0.7152 * int(bg_hex[2:4], 16)
                      + 0.0722 * int(bg_hex[4:6], 16)) / 255.0
            if abs(lum_logo - lum_bg) < 0.32:
                pad = 0.16
                chip = "FFFFFF" if lum_logo < 0.5 else "111111"
                d.card(s, x - pad, y - pad, w_in + 2 * pad, height + 2 * pad,
                       fill=chip, radius=0.16)
        s.shapes.add_picture(logo, Inches(x), Inches(y),
                             height=Inches(height), width=Inches(w_in))
    except Exception:
        pass


def slide_cover(d, data, logo):
    meta, p = d.meta, d.p
    s = d.slide(dark=True)
    ink, muted = p["invert_ink"], p["invert_muted"]

    y = 2.05
    if logo and os.path.exists(logo):
        _place_logo(d, s, logo, M, 1.05, p["invert_bg"])

    d.text(s, M, y, 9.4, 0.6, meta["shop_name"], size=15, color=muted,
           font=p["font_body"])
    d.text(s, M, y + 0.52, 10.8, 1.5, T("cover.title"), size=48,
           bold=True, color=ink, font=p["font_heading"])
    d.text(s, M, y + 1.62, 10.8, 0.9, meta["period_label"], size=34,
           color=p["invert_accent"], font=p["font_heading"], bold=True)

    base = meta.get("comparison_sentence")
    if base:
        d.text(s, M, y + 2.7, 10.8, 0.5, base, size=11.5, color=muted)

    d.text(s, M, FOOT, 8.0, 0.3,
           T("cover.generated", date=meta["generated_on"]),
           size=9, color=muted, anchor=MSO_ANCHOR.MIDDLE)
    if meta.get("agency"):
        d.text(s, W - M - 3.4, FOOT, 3.4, 0.3,
               T("cover.agency", agency=meta["agency"]), size=9, color=muted,
               align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    d.page += 1


def slide_toc(d, chapters, data):
    """Sommaire. Sur un rapport de 20 slides et plus, il évite que le lecteur
    se perde et lui permet d'aller directement au chapitre qui l'intéresse."""
    p = d.p
    s = d.slide()
    d.frame(s, T("toc.title"), d.meta.get("kicker"))
    entries = [(T("toc.essentials"), T("toc.essentials_blurb"))]
    if data.get("priorities"):
        entries.append((T("toc.plan"), T("toc.plan_blurb")))
    entries += [(c["title"], c["blurb"]) for c in chapters]
    entries.append((T("toc.method"), T("toc.method_blurb")))

    n = len(entries)
    col = (n + 1) // 2
    half = (W - 2 * M - 0.6) / 2
    rh = min(0.78, 4.5 / max(1, col))
    for i, (title, blurb) in enumerate(entries):
        ci, ri = i // col, i % col
        x = M + ci * (half + 0.6)
        y = BODY_TOP + ri * rh
        d.text(s, x, y, 0.42, rh - 0.1, f"{i + 1:02d}", size=13, bold=True,
               color=p["accent_line"], font=p["font_heading"],
               anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, x + 0.46, y, half - 0.46, rh - 0.1,
               [{"t": title, "size": 13, "bold": True, "color": p["ink"]},
                {"t": blurb, "size": 9.5, "color": p["muted"], "space_before": 2}],
               anchor=MSO_ANCHOR.MIDDLE)


def slide_kpis(d, data):
    p, meta = d.p, d.meta
    kpis = (data.get("kpis") or [])[:4]
    if not kpis:
        return
    s = d.slide()
    cy = d.frame(s, T("slides.kpis"), meta.get("kicker"),
                 hint=lecture(data, "kpis")) + 0.06

    n = len(kpis)
    gap = 0.28
    cw = min(3.2, (W - 2 * M - gap * (n - 1)) / n) if n < 4 \
        else (W - 2 * M - gap * (n - 1)) / n
    ch = 2.3

    for i, k in enumerate(kpis):
        x = M + i * (cw + gap)
        d.card(s, x, cy, cw, ch)
        pad = 0.34
        d.text(s, x + pad, cy + 0.30, cw - 2 * pad, 0.32, k["label"].upper(),
               size=9, bold=True, color=p["muted"])
        d.text(s, x + pad, cy + 0.70, cw - 2 * pad, 0.78,
               fmt(k.get("value"), k.get("format", "number"), meta["currency"]),
               size=31, bold=True, color=p["ink"], font=p["font_heading"])
        ly = cy + 1.54
        for ref_key, ref_label in meta["comparison_labels"].items():
            txt, sense = delta(k.get("value"), (k.get("compare") or {}).get(ref_key))
            if txt is None:
                continue
            col = p["up"] if sense > 0 else (p["down"] if sense < 0 else p["muted"])
            d.text(s, x + pad, ly, 1.05, 0.26, txt, size=11.5, bold=True, color=col)
            d.text(s, x + pad + 1.05, ly, cw - 2 * pad - 1.05, 0.26,
                   T("common.vs", label=ref_label), size=9, color=p["muted"],
                   anchor=MSO_ANCHOR.MIDDLE)
            ly += 0.3

    notes = data.get("essentials_notes") or []
    if notes:
        ny = cy + ch + 0.34
        d.card(s, M, ny, W - 2 * M, 1.14, fill=p["surface_alt"])
        runs = [{"t": T("labels.facts"), "size": 9, "bold": True, "color": p["muted"]}]
        for t in notes[:3]:
            runs.append({"t": t, "size": 11.5, "color": p["ink"], "space_before": 5})
        d.text(s, M + 0.34, ny + 0.22, W - 2 * M - 0.68, 0.8, runs)


def slide_highlights(d, data):
    """Les constats de saillance. C'est la slide qui mâche le travail : elle
    dit ce qui sort de l'ordinaire, en chiffres, sans jamais dire pourquoi."""
    hl = data.get("highlights") or []
    if not hl:
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.highlights"), d.meta.get("kicker"),
                 hint=lecture(data, "highlights"))
    hl = hl[:6]
    n = len(hl)
    col = (n + 1) // 2 if n > 3 else n
    half = (W - 2 * M - 0.4) / 2 if n > 3 else (W - 2 * M)
    rh = min(1.28, (4.55 - 0.1) / max(1, col))
    for i, h in enumerate(hl):
        ci, ri = (i // col, i % col) if n > 3 else (0, i)
        x = M + ci * (half + 0.4)
        y = y0 + ri * rh
        d.card(s, x, y, half, rh - 0.16)
        d.text(s, x + 0.30, y + 0.20, half - 0.60, 0.26, h["tag"].upper(),
               size=8.5, bold=True, color=p["accent_line"])
        d.text(s, x + 0.30, y + 0.50, half - 0.60, rh - 0.76, h["text"],
               size=12, color=p["ink"], spacing=1.15)


def slide_plan(d, data):
    """
    Le plan des 90 prochains jours. C'est la seule slide qui regarde devant,
    et elle ne tient que parce que chaque objectif est un niveau que la
    boutique a déjà atteint, avec un gain calculé au lieu d'être décrété.
    """
    prios = data.get("priorities") or []
    if not prios:
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.plan"), d.meta.get("kicker"),
                 hint=lecture(data, "plan"))

    total_w = W - 2 * M
    gauche = total_w * 0.60
    droite = total_w - gauche - 0.3
    n = len(prios)
    ch = min(1.46, (FOOT - 0.72 - y0) / max(1, n) - 0.14)

    for i, c in enumerate(prios):
        y = y0 + i * (ch + 0.14)
        d.card(s, M, y, total_w, ch)
        d.text(s, M + 0.32, y + 0.16, 2.0, 0.24, T("labels.priority", n=c["rank"]),
               size=8.5, bold=True, color=p["accent_line"])
        d.text(s, M + 0.32, y + 0.42, gauche - 0.5, 0.34, c["title"], size=15,
               bold=True, color=p["ink"], font=p["font_heading"])
        d.text(s, M + 0.32, y + 0.82, gauche - 0.5, ch - 0.94, c["finding"],
               size=9.5, color=p["muted"], spacing=1.14)

        x1 = M + gauche + 0.3
        d.text(s, x1, y + 0.16, droite - 0.32, 0.24, T("labels.goal_90"),
               size=8.5, bold=True, color=p["muted"])
        d.text(s, x1, y + 0.42, droite - 0.32, 0.34, c["goal"], size=12,
               bold=True, color=p["ink"])
        d.text(s, x1, y + 0.80, droite - 0.32, 0.26,
               T("labels.gain_line", label=c["gain_label"], value=c["gain_txt"]),
               size=10, bold=True, color=p["ink"])
        d.text(s, x1, y + 1.06, droite - 0.32, 0.26,
               T("labels.milestones", values="  ·  ".join(c["milestones_txt"])),
               size=9, color=p["muted"])

    d.text(s, M, FOOT - 0.62, total_w, 0.55, T("captions.plan"),
           size=9, color=p["muted"])


def slide_suivi(d, data):
    """Où en est le plan du mois précédent. Sans jugement : le jalon visé, la
    valeur mesurée, et si le jalon est tenu."""
    lignes = data.get("plan_suivi") or []
    if not lignes:
        return
    p = d.p
    s = d.slide()
    y = d.frame(s, T("slides.follow_up"), d.meta.get("kicker"),
                hint=lecture(data, "follow_up")) + 0.1

    cols = [T("tables.project"), T("tables.starting_point"),
            T("tables.milestone"), T("tables.measured"),
            T("tables.milestone_met")]
    widths = [0.30, 0.17, 0.17, 0.20, 0.16]
    total_w = W - 2 * M
    aligns = [PP_ALIGN.LEFT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT, PP_ALIGN.RIGHT,
              PP_ALIGN.RIGHT]

    x = M
    for i, col in enumerate(cols):
        cw = total_w * widths[i]
        d.text(s, x + 0.12, y, cw - 0.24, 0.32, col.upper(), size=8.5, bold=True,
               color=p["muted"], align=aligns[i], anchor=MSO_ANCHOR.MIDDLE)
        x += cw
    y += 0.38
    d.bar(s, M, y - 0.06, total_w, 0.01, p["hairline"])

    rh = 0.46
    for ri, l in enumerate(lignes[:6]):
        if ri % 2 == 1:
            d.bar(s, M, y, total_w, rh, p["surface"])
        cells = [l["title"], l["start"], l["milestone"], l["measured"],
                 T("common.yes") if l["met"] else T("common.no")]
        x = M
        for ci, cell in enumerate(cells):
            cw = total_w * widths[ci]
            col = p["ink"]
            if ci == 4:
                col = p["up"] if l["met"] else p["down"]
            d.text(s, x + 0.12, y, cw - 0.24, rh, cell, size=10.5,
                   bold=(ci == 4), color=col, align=aligns[ci],
                   anchor=MSO_ANCHOR.MIDDLE)
            x += cw
        y += rh

    d.text(s, M, FOOT - 0.62, total_w, 0.55, T("captions.follow_up"),
           size=9, color=p["muted"])


def slide_trend(d, data):
    ts = data.get("timeseries")
    if not ts or not ts.get("labels"):
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.trend"), d.meta.get("kicker"),
                 hint=lecture(data, "trend"))
    series = [(x["name"], x["values"]) for x in ts["series"]]
    d.line_chart(s, M - 0.12, y0, W - 2 * M + 0.24, 4.05, ts["labels"], series)
    if ts.get("caption"):
        d.text(s, M, y0 + 4.15, W - 2 * M, 0.5, ts["caption"], size=9,
               color=p["muted"])


def slide_seasonality(d, data):
    """Le mois replacé dans son histoire. Sans cette slide, comparer août à
    juillet n'a aucun sens sur une activité saisonnière."""
    hist = data.get("seasonality") or []
    if len(hist) < 6:
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.seasonality"), d.meta.get("kicker"),
                 hint=lecture(data, "seasonality"))
    hist = hist[-24:]
    cats = [h["label"] for h in hist]
    vals = [h.get("value") or 0 for h in hist]
    # Le mois courant prend l'accent, les autres restent neutres.
    cols = [p["series"][5]] * (len(vals) - 1) + [p["series"][0]]
    d.bar_chart(s, M - 0.1, y0, W - 2 * M + 0.2, 3.35, cats, vals,
                value_fmt=d.money_fmt, labels=False, colors=cols, gap=35)

    cur = vals[-1]
    douze = vals[-12:]
    rang = sorted(douze, reverse=True).index(cur) + 1 if cur in douze else None
    moy = sum(douze) / len(douze) if douze else 0
    rows = [
        (T("labels.month_revenue"), money(cur, d.meta["currency"], True)),
        (T("labels.twelve_month_avg"), money(moy, d.meta["currency"], True)),
    ]
    if rang:
        rows.append((T("labels.rank_twelve"),
                     T("labels.rank_value", n=rang, total=len(douze))))
    _rows(d, s, M, y0 + 3.55, W - 2 * M, rows, rh=0.34)


def slide_acquisition(d, data):
    funnel = data.get("funnel")
    sources = data.get("traffic_sources") or []
    if not funnel and not sources:
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.acquisition"), d.meta.get("kicker"),
                 hint=lecture(data, "acquisition"))
    half = (W - 2 * M - 0.42) / 2

    if funnel:
        d.text(s, M, y0 - 0.04, half, 0.3, T("labels.purchase_path"), size=9,
               bold=True, color=p["muted"])
        steps = [(T("funnel.sessions"), funnel.get("sessions")),
                 (T("funnel.cart"), funnel.get("cart")),
                 (T("funnel.checkout"), funnel.get("checkout")),
                 (T("funnel.purchase"), funnel.get("purchase"))]
        steps = [st for st in steps if st[1] is not None]
        top = max((v for _, v in steps), default=1) or 1
        y = y0 + 0.42
        rh, rgap = 0.56, 0.34
        for i, (label, val) in enumerate(steps):
            frac = max(0.004, val / top)
            d.bar(s, M, y, half, rh, p["surface"], radius=0.14)
            d.bar(s, M, y, max(0.06, half * frac), rh, p["series"][0], radius=0.14)
            d.text(s, M + 0.34, y, half * 0.55, rh, label, size=11.5, bold=True,
                   color=p["ink"], anchor=MSO_ANCHOR.MIDDLE)
            d.text(s, M + half * 0.55, y, half * 0.45 - 0.24, rh,
                   f"{integer(val)}   ·   {percent(val / top)}", size=11,
                   color=p["ink"], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            if i < len(steps) - 1 and steps[i][1]:
                rate = steps[i + 1][1] / steps[i][1]
                d.text(s, M + 0.34, y + rh + 0.02, half - 0.34, rgap - 0.04,
                       T("funnel.step_rate", rate=percent(rate)), size=9,
                       color=p["muted"], anchor=MSO_ANCHOR.MIDDLE)
            y += rh + rgap
        d.text(s, M, y - rgap + 0.12, half, 0.4, T("funnel.caption"),
               size=9, color=p["muted"])

    if sources:
        x0 = M + half + 0.42
        d.text(s, x0, y0 - 0.04, half, 0.3, T("labels.revenue_by_source"),
               size=9, bold=True, color=p["muted"])
        top_src = sources[:6]
        d.bar_chart(s, x0 - 0.1, y0 + 0.3, half + 0.2, 3.6,
                    [t["label"] for t in top_src][::-1],
                    [t.get("sales") or 0 for t in top_src][::-1],
                    horizontal=True, value_fmt=d.money_fmt, gap=55)
        d.text(s, x0, y0 + 4.0, half, 0.5, T("captions.sources"), size=9,
               color=p["muted"])


def slide_customers(d, data):
    cust = data.get("customers")
    if not cust:
        return
    p = d.p
    s = d.slide()
    y0 = d.frame(s, T("slides.customers"), d.meta.get("kicker"),
                 hint=lecture(data, "customers"))
    half = (W - 2 * M - 0.42) / 2

    d.card(s, M, y0, half, 4.0)
    d.text(s, M + 0.34, y0 + 0.26, half - 0.68, 0.3, T("labels.buyer_split"),
           size=9, bold=True, color=p["muted"])
    new, ret = cust.get("new") or 0, cust.get("returning") or 0
    if new or ret:
        d.doughnut(s, M + 0.2, y0 + 0.58, half - 0.4, 2.2,
                   [T("labels.new_customers"), T("labels.returning_customers")],
                   [new, ret])
    rows = [(T("labels.customers_ordered"), integer(cust.get("total"))),
            (T("labels.returning_rate"), percent(cust.get("returning_rate")))]
    if cust.get("repeat_rate") is not None:
        rows.append((T("labels.repeat_rate"), percent(cust["repeat_rate"])))
    if cust.get("days_to_second_order") is not None:
        rows.append((T("labels.days_to_second"),
                     T("labels.days", n=integer(cust["days_to_second_order"]))))
    _rows(d, s, M + 0.34, y0 + 2.86, half - 0.68, rows, rh=0.38)

    x0 = M + half + 0.42
    ch = min(4.0, 0.74 + 3 * 0.42 + 0.36)
    d.card(s, x0, y0, half, ch)
    d.text(s, x0 + 0.34, y0 + 0.26, half - 0.68, 0.3, T("labels.volumes"),
           size=9, bold=True, color=p["muted"])
    _rows(d, s, x0 + 0.34, y0 + 0.72, half - 0.68,
          [(T("labels.new_customers"), integer(new)),
           (T("labels.returning_customers"), integer(ret)),
           (T("labels.total"), integer(cust.get("total")))], rh=0.42)
    note = cust.get("note")
    if note:
        d.text(s, x0, y0 + ch + 0.22, half, 0.9, note, size=9.5, color=p["muted"],
               spacing=1.15)


def slide_margin(d, data):
    mg = data.get("margin")
    if not mg:
        return
    p, cur = d.p, d.meta["currency"]
    s = d.slide()
    y0 = d.frame(s, T("slides.margin"), d.meta.get("kicker"),
                 hint=lecture(data, "margin"))
    half = (W - 2 * M - 0.42) / 2

    d.card(s, M, y0, half, 2.6)
    d.text(s, M + 0.34, y0 + 0.28, half - 0.68, 0.3, T("labels.gross_margin"),
           size=9, bold=True, color=p["muted"])
    d.text(s, M + 0.34, y0 + 0.66, half - 0.68, 0.85,
           money(mg.get("gross_profit"), cur, True), size=34, bold=True,
           color=p["ink"], font=p["font_heading"])
    d.text(s, M + 0.34, y0 + 1.56, half - 0.68, 0.3,
           T("labels.margin_of_net", rate=percent(mg.get("margin_rate"))),
           size=11.5, color=p["muted"])

    x0 = M + half + 0.42
    rows = [(T("labels.net_sales"), money(mg.get("net_sales"), cur, True)),
            (T("labels.cogs"), money(mg.get("cogs"), cur, True)),
            (T("labels.gross_profit"), money(mg.get("gross_profit"), cur, True)),
            (T("labels.margin_rate"), percent(mg.get("margin_rate")))]
    ref_rate = mg.get("compare_margin_rate")
    if ref_rate is not None:
        key = next(iter(d.meta.get("comparison_labels") or {}), None)
        ref_label = (d.meta.get("comparison_labels") or {}).get(
            key, T("common.reference_period"))
        rows.append((T("labels.margin_rate_ref", label=ref_label),
                     percent(ref_rate)))
    d.card(s, x0, y0, half, 0.74 + len(rows) * 0.42 + 0.36)
    d.text(s, x0 + 0.34, y0 + 0.28, half - 0.68, 0.3, T("labels.detail"),
           size=9, bold=True, color=p["muted"])
    _rows(d, s, x0 + 0.34, y0 + 0.74, half - 0.68, rows)

    notes = []
    if mg.get("incomplete"):
        notes.append(mg["incomplete"])
    notes.append(T("captions.margin"))
    d.text(s, M, y0 + 3.0, W - 2 * M, 1.0,
           [{"t": t, "size": 9.5, "color": p["muted"],
             "space_before": 4 if i else 0} for i, t in enumerate(notes)],
           spacing=1.2)


def _rows(d, s, x, y, w, rows, rh=0.42):
    p = d.p
    for i, (label, value) in enumerate(rows):
        yy = y + i * rh
        d.text(s, x, yy, w * 0.62, rh, label, size=11, color=p["muted"],
               anchor=MSO_ANCHOR.MIDDLE)
        d.text(s, x + w * 0.62, yy, w * 0.38, rh, value, size=11.5, bold=True,
               color=p["ink"], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(rows) - 1:
            d.bar(s, x, yy + rh - 0.01, w, 0.008, p["hairline"])


def slide_divider(d, title, subtitle=None):
    p = d.p
    s = d.slide(dark=True)
    d.text(s, M, 3.0, W - 2 * M, 1.0, title, size=40, bold=True,
           color=p["invert_ink"], font=p["font_heading"])
    if subtitle:
        d.text(s, M, 4.05, W - 2 * M, 0.6, subtitle, size=13,
               color=p["invert_muted"])
    d.footer(s, dark=True)


def slide_table(d, title, columns, rows, caption=None, hint=None, aligns=None,
                highlight_last=False, max_rows=11):
    if not rows:
        return
    p = d.p
    s = d.slide()
    y = d.frame(s, title, d.meta.get("kicker"), hint=hint) + 0.08
    n = len(columns)
    aligns = aligns or ([PP_ALIGN.LEFT] + [PP_ALIGN.RIGHT] * (n - 1))
    widths = [0.40] + [0.60 / (n - 1)] * (n - 1) if n > 1 else [1.0]
    total_w = W - 2 * M
    rows = rows[:max_rows]
    # La hauteur de ligne doit réserver la place de la légende, sinon un
    # tableau de 10 lignes la fait passer sous la dernière ligne.
    cap_h = 0.62 if caption else 0.10
    avail = FOOT - 0.15 - cap_h - (y + 0.38)
    rh = max(0.26, min(0.42, avail / max(1, len(rows))))

    x = M
    for i, col in enumerate(columns):
        cw = total_w * widths[i]
        d.text(s, x + 0.12, y, cw - 0.24, 0.32, col.upper(), size=8.5, bold=True,
               color=p["muted"], align=aligns[i], anchor=MSO_ANCHOR.MIDDLE)
        x += cw
    y += 0.38
    d.bar(s, M, y - 0.06, total_w, 0.01, p["hairline"])

    for ri, row in enumerate(rows):
        last = highlight_last and ri == len(rows) - 1
        if ri % 2 == 1 and not last:
            d.bar(s, M, y, total_w, rh, p["surface"])
        if last:
            d.bar(s, M, y, total_w, rh, p["surface_alt"])
        x = M
        for ci, cell in enumerate(row):
            cw = total_w * widths[ci]
            d.text(s, x + 0.12, y, cw - 0.24, rh, str(cell), size=10.5, bold=last,
                   color=p["ink"], align=aligns[ci], anchor=MSO_ANCHOR.MIDDLE)
            x += cw
        y += rh
    if caption:
        d.text(s, M, max(y + 0.16, FOOT - 0.62), total_w, 0.55, caption, size=9,
               color=p["muted"])


def slide_chart_and_rows(d, title, cats, vals, rows, chart_title=None,
                         rows_title=None, caption=None, value_fmt=None,
                         hint=None):
    p = d.p
    value_fmt = value_fmt or d.money_fmt
    s = d.slide()
    y0 = d.frame(s, title, d.meta.get("kicker"), hint=hint)
    half = (W - 2 * M - 0.42) / 2
    if cats:
        if chart_title:
            d.text(s, M, y0 - 0.04, half, 0.3, chart_title.upper(), size=9,
                   bold=True, color=p["muted"])
        d.bar_chart(s, M - 0.1, y0 + 0.3, half + 0.2, 3.6, cats, vals,
                    horizontal=True, value_fmt=value_fmt, gap=55)
    if rows:
        x0 = M + half + 0.42
        ch = min(4.0, 0.74 + len(rows) * 0.42 + 0.36)
        d.card(s, x0, y0, half, ch)
        if rows_title:
            d.text(s, x0 + 0.34, y0 + 0.28, half - 0.68, 0.3, rows_title.upper(),
                   size=9, bold=True, color=p["muted"])
        _rows(d, s, x0 + 0.34, y0 + 0.74, half - 0.68, rows)
    if caption:
        d.text(s, M, FOOT - 0.62, W - 2 * M, 0.55, caption, size=9, color=p["muted"])


def slide_text_blocks(d, title, blocks, hint=None):
    p = d.p
    s = d.slide()
    y0 = d.frame(s, title, d.meta.get("kicker"), hint=hint)
    per_col = (len(blocks) + 1) // 2
    cols = [blocks[:per_col], blocks[per_col:]]
    half = (W - 2 * M - 0.5) / 2
    cpl = 78

    def slot(b):
        lines = max(1, -(-len(b["text"]) // cpl))
        return 0.30 + lines * 0.185 + 0.22

    heights = [max(slot(cols[0][i]) if i < len(cols[0]) else 0,
                   slot(cols[1][i]) if i < len(cols[1]) else 0)
               for i in range(per_col)]
    for ci, col in enumerate(cols):
        x = M + ci * (half + 0.5)
        y = y0
        for i, b in enumerate(col):
            d.text(s, x, y, half, heights[i],
                   [{"t": b["title"], "size": 11.5, "bold": True, "color": p["ink"]},
                    {"t": b["text"], "size": 10, "color": p["muted"],
                     "space_before": 3}], spacing=1.12)
            y += heights[i]


# ---------------------------------------------------------------------------
# composition du rapport
# ---------------------------------------------------------------------------

def _parts(items, key, total=None):
    """Adds the relative share to a list of {label, <key>}."""
    tot = total if total else sum(abs(i.get(key) or 0) for i in items)
    return [(i, (abs(i.get(key) or 0) / tot) if tot else 0) for i in items]


def _gen_sales_detail(d, data):
    cur = d.meta["currency"]
    sd = data["sales_detail"]
    order = [(k, T(f"sales_detail.{k}")) for k in
             ("gross_sales", "discounts", "sales_reversals", "net_sales",
              "shipping_charges", "taxes", "total_sales")]
    magnitude = {"discounts", "sales_reversals"}
    rows = []
    for key, label in order:
        if sd.get(key) is None:
            continue
        txt, _ = delta(sd[key], (sd.get("compare") or {}).get(key),
                       magnitude=key in magnitude)
        rows.append([label, money(sd[key], cur, True), txt or "—"])
    slide_table(d, T("slides.sales_detail"),
                [T("tables.item"), T("tables.amount"), T("tables.change")], rows,
                hint=lecture(data, "sales_detail"),
                caption=T("captions.sales_detail"), highlight_last=True)


def _gen_channels(d, data):
    cur = d.meta["currency"]
    ch = data["channels"]
    tot = sum(c.get("sales") or 0 for c in ch)
    rows = [(c["label"], f"{money(c.get('sales'), cur, True)}  ·  "
                         f"{percent((c.get('sales') or 0) / tot if tot else 0)}")
            for c in ch[:6]]
    slide_chart_and_rows(
        d, T("slides.channels"), [c["label"] for c in ch[:6]][::-1],
        [c.get("sales") or 0 for c in ch[:6]][::-1], rows,
        chart_title=T("labels.revenue_by_channel"), rows_title=T("labels.split"),
        hint=lecture(data, "channels"), caption=T("captions.channels"))


def _gen_categories(d, data):
    cur = d.meta["currency"]
    cats = [c for c in data["categories"] if c.get("label")]
    tot = sum(c.get("sales") or 0 for c in cats)
    rows = [(c["label"], f"{money(c.get('sales'), cur, True)}  ·  "
                         f"{percent((c.get('sales') or 0) / tot if tot else 0)}")
            for c in cats[:6]]
    slide_chart_and_rows(
        d, T("slides.categories"), [c["label"] for c in cats[:6]][::-1],
        [c.get("sales") or 0 for c in cats[:6]][::-1], rows,
        chart_title=T("labels.gross_by_category"), rows_title=T("labels.split"),
        hint=lecture(data, "categories"), caption=T("captions.categories"))


def _gen_variants(d, data):
    cur = d.meta["currency"]
    vs = data["variants"]
    rows = [[v["label"], money(v.get("sales"), cur, True)] for v in vs[:11]]
    slide_table(d, T("slides.variants"),
                [T("tables.variant"), T("tables.gross_sales")], rows,
                hint=lecture(data, "variants"), caption=T("captions.variants"))


def _gen_sales_country(d, data):
    cur = d.meta["currency"]
    cs = data["sales_by_country"]
    tot = sum(c.get("sales") or 0 for c in cs)
    rows = [(c["label"], f"{money(c.get('sales'), cur, True)}  ·  "
                         f"{percent((c.get('sales') or 0) / tot if tot else 0)}")
            for c in cs[:6]]
    slide_chart_and_rows(
        d, T("slides.sales_country"), [c["label"] for c in cs[:6]][::-1],
        [c.get("sales") or 0 for c in cs[:6]][::-1], rows,
        chart_title=T("labels.revenue_by_country"), rows_title=T("labels.split"),
        hint=lecture(data, "sales_country"), caption=T("captions.sales_country"))


def _gen_top_products(d, data):
    cur = d.meta["currency"]
    tp = data["top_products"]
    gross = (data.get("sales_detail") or {}).get("gross_sales")
    tot = gross or sum(p.get("sales") or 0 for p in tp)
    rows = [[p["label"], money(p.get("sales"), cur, True), integer(p.get("orders")),
             percent((p.get("sales") or 0) / tot if tot else 0)] for p in tp[:11]]
    slide_table(d, T("slides.top_products"),
                [T("tables.product"), T("tables.gross_sales"), T("tables.orders"),
                 T("tables.share_of_sales")], rows,
                hint=lecture(data, "top_products"),
                caption=T("captions.top_products"))


def _gen_inventory(d, data):
    inv = data["inventory"]
    rows = [[i["label"], integer(i.get("sold")), integer(i.get("ending_units")),
             percent(i.get("sell_through"))] for i in inv[:11]]
    slide_table(d, T("slides.inventory"),
                [T("tables.product"), T("tables.units_sold"),
                 T("tables.ending_stock"), T("tables.sell_through")], rows,
                hint=lecture(data, "inventory"), caption=T("captions.inventory"))


def _gen_returns(d, data):
    cur = d.meta["currency"]
    rt = data["returns"]
    rows = [(T("sales_detail.sales_reversals"), money(rt.get("amount"), cur, True)),
            (T("tables.share_of_sales"), percent(rt.get("rate")))]
    if rt.get("orders") is not None:
        rows.append((T("tables.orders"), integer(rt["orders"])))
    top = rt.get("top") or []
    slide_chart_and_rows(
        d, T("slides.returns"),
        [t["label"] for t in top[:6]][::-1],
        [abs(t.get("amount") or 0) for t in top[:6]][::-1], rows,
        chart_title=T("labels.most_returned"), rows_title=T("labels.overview"),
        hint=lecture(data, "returns"), caption=T("captions.returns"))


def _gen_traffic(d, data):
    dev, ctry = data.get("devices"), data.get("countries")
    cats = [x["label"] for x in (ctry or [])[:6]][::-1]
    vals = [x.get("sessions") or 0 for x in (ctry or [])[:6]][::-1]
    rows = [(x["label"], f"{integer(x.get('sessions'))}  ·  "
                         f"{percent(x.get('conversion_rate'))}")
            for x in (dev or [])[:5]]
    slide_chart_and_rows(
        d, T("slides.traffic"), cats, vals, rows,
        chart_title=T("labels.sessions_by_country"),
        rows_title=T("labels.sessions_by_device"), value_fmt="#,##0",
        hint=lecture(data, "traffic"), caption=T("captions.traffic"))


def _gen_extra(key, title):
    def gen(d, data):
        blk = data[key]
        rows = [(r["label"], r["value"]) for r in blk.get("rows", [])]
        slide_chart_and_rows(
            d, title, [c["label"] for c in blk.get("chart", [])][::-1],
            [c["value"] for c in blk.get("chart", [])][::-1], rows,
            chart_title=blk.get("chart_title", T("labels.attributed_revenue")),
            rows_title=T("labels.overview"), hint=lecture(data, key),
            caption=blk.get("caption"))
    return gen


def plan(data):
    """
    Compose le rapport à partir du profil de la boutique.

    Le profil ne change pas les chiffres, il décide des chapitres présents et
    de leur ordre. Une boutique qui vit de l'acquisition ne lit pas son rapport
    dans le même ordre qu'une boutique qui vit du réachat.
    """
    prof = data.get("profile") or {}
    chapters = []

    ventes = []
    if data.get("timeseries"):
        ventes.append(slide_trend)
    if len(data.get("seasonality") or []) >= 6:
        ventes.append(slide_seasonality)
    if data.get("sales_detail"):
        ventes.append(_gen_sales_detail)
    if len(data.get("channels") or []) >= 2:
        ventes.append(_gen_channels)
    if ventes:
        chapters.append({"key": "sales", "title": T("chapters.sales"),
                         "blurb": T("chapters.sales_blurb"),
                         "gens": ventes})

    acq = []
    if data.get("funnel") or data.get("traffic_sources"):
        acq.append(slide_acquisition)
    if data.get("devices") or data.get("countries"):
        acq.append(_gen_traffic)
    if len(data.get("sales_by_country") or []) >= 2:
        acq.append(_gen_sales_country)
    if acq:
        chapters.append({"key": "acquisition", "title": T("chapters.acquisition"),
                         "blurb": T("chapters.acquisition_blurb"),
                         "gens": acq})

    prod = []
    if data.get("top_products"):
        prod.append(_gen_top_products)
    if len([c for c in (data.get("categories") or []) if c.get("label")]) >= 2:
        prod.append(_gen_categories)
    if len(data.get("variants") or []) >= 3:
        prod.append(_gen_variants)
    if data.get("inventory"):
        prod.append(_gen_inventory)
    if data.get("returns"):
        prod.append(_gen_returns)
    if prod:
        chapters.append({"key": "products", "title": T("chapters.products"),
                         "blurb": T("chapters.products_blurb"),
                         "gens": prod})

    clients = []
    if data.get("customers"):
        clients.append(slide_customers)
    if clients:
        chapters.append({"key": "customers", "title": T("chapters.customers"),
                         "blurb": T("chapters.customers_blurb"),
                         "gens": clients})

    renta = []
    if data.get("margin"):
        renta.append(slide_margin)
    if renta:
        chapters.append({"key": "profitability", "title": T("chapters.profitability"),
                         "blurb": T("chapters.profitability_blurb"),
                         "gens": renta})

    mkt = []
    if data.get("klaviyo"):
        mkt.append(_gen_extra("klaviyo", T("slides.klaviyo")))
    if data.get("ads"):
        mkt.append(_gen_extra("ads", T("slides.ads")))
    if mkt:
        chapters.append({"key": "marketing", "title": T("chapters.marketing"),
                         "blurb": T("chapters.marketing_blurb"),
                         "gens": mkt})

    # Le profil réordonne. Une activité de réachat met ses clients en avant ;
    # une activité d'achat unique met son acquisition en avant.
    # `driver` is the English key; `moteur` is kept so contexts written by
    # earlier French-only versions keep working.
    moteur = prof.get("driver") or prof.get("moteur")
    ordre = {"retention": ["sales", "customers", "products", "acquisition",
                           "profitability", "marketing"],
             "acquisition": ["sales", "acquisition", "products", "customers",
                             "profitability", "marketing"]}.get(moteur)
    if ordre:
        chapters.sort(key=lambda c: ordre.index(c["key"])
                      if c["key"] in ordre else 99)
    return chapters


def build(data, palette, logo, out):
    meta = data["meta"]
    meta.setdefault("currency", "EUR")
    meta.setdefault("agency", None)
    meta["kicker"] = meta.get("kicker") or meta.get("comparison_sentence")
    meta.setdefault("comparison_labels", {})

    if not data.get("highlights"):
        try:
            import highlights as hl_mod
            data["highlights"] = hl_mod.compute(data, lang=_L)
        except Exception:
            data["highlights"] = []

    prio_state = None
    if not data.get("priorities"):
        try:
            import priorities as pr_mod
            data["priorities"] = pr_mod.compute(data, lang=_L)
            prio_state = pr_mod.to_state(data["priorities"], meta)
            if data.get("previous_plan"):
                data["plan_suivi"] = pr_mod.follow_up(
                    data["previous_plan"], data, lang=_L)
        except Exception:
            data["priorities"] = []

    # Un plan sans sa note de méthode n'est pas défendable : on garantit sa
    # présence plutôt que de compter sur le rédacteur.
    if data.get("priorities"):
        blocs = data.setdefault("methodology", [])
        if not any("plan" in (b.get("title") or "").lower() for b in blocs):
            blocs.append({"title": T("methodology.plan_title"),
                          "text": T("methodology.plan_text")})

    d = Deck(palette, meta)
    chapters = plan(data)

    slide_cover(d, data, logo)
    if sum(len(c["gens"]) for c in chapters) >= 8:
        slide_toc(d, chapters, data)
    slide_kpis(d, data)
    slide_highlights(d, data)
    slide_suivi(d, data)
    slide_plan(d, data)

    for c in chapters:
        d.chapter = c["title"]
        slide_divider(d, c["title"], c["blurb"])
        for gen in c["gens"]:
            gen(d, data)
    d.chapter = None

    blocks = data.get("methodology") or []
    if blocks:
        slide_text_blocks(d, T("slides.methodology"), blocks,
                          hint=lecture(data, "methodology"))

    d.prs.save(out)
    return out, [c["title"] for c in chapters], data.get("highlights") or [], prio_state


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", required=True)
    ap.add_argument("--theme", help="config/settings_data.json du theme")
    ap.add_argument("--config", help="surcharges client (json)")
    ap.add_argument("--logo")
    ap.add_argument("-o", "--out", default="rapport.pptx")
    ap.add_argument("--state", help="file to write the plan to, to be stored "
                                    "back in the shop metafield")
    ap.add_argument("--lang", help="report language (default: meta.lang, else en)")
    a = ap.parse_args()

    data = json.load(open(a.data, encoding="utf-8"))
    code = a.lang or (data.get("meta") or {}).get("lang") or "en"
    lang = Lang(code)
    set_lang(lang)
    if lang.code != code:
        print(f"warning: language {code!r} not available, falling back to "
              f"{lang.code!r}. Available: {', '.join(available())}", file=sys.stderr)
    raw = open(a.theme, encoding="utf-8").read() if a.theme and os.path.exists(a.theme) else ""
    over = json.load(open(a.config, encoding="utf-8")) if a.config and os.path.exists(a.config) else {}
    palette = build_palette(raw, over.get("brand") if "brand" in over else over)
    if over.get("agency") is not None:
        data.setdefault("meta", {})["agency"] = over["agency"]

    out, chapters, hl, prio_state = build(data, palette, logo=a.logo, out=a.out)
    if a.state and prio_state:
        with open(a.state, "w", encoding="utf-8") as f:
            json.dump(prio_state, f, ensure_ascii=False, indent=1)
    print(json.dumps({
        "out": out,
        "lang": lang.code,
        "palette_source": palette["source"],
        "bg": palette["bg"],
        "accent": palette["accent"],
        "chapters": chapters,
        "findings": [h["text"] for h in hl],
        "plan": [{"rank": c["rank"], "title": c["title"],
                  "goal": c["goal"], "gain": c["gain_txt"]}
                 for c in (data.get("priorities") or [])],
        "state_file": a.state if (a.state and prio_state) else None,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
